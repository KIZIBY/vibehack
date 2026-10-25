"""Работы участников: создание одной и пачкой, список, детали.

Файлы пишутся в ``projects/<pid>/submissions/<sid>/raw/<material_id>/<имя>`` (agent-run.md §1), строки БД — одной
транзакцией; любой сбой между ними убирает уже записанные файлы. Уникальность названия без учёта регистра держится
на транзакции: первым оператором она обновляет строку проекта (блокировка записи), затем сверяет названия и вставляет;
уникальный ключ БД ``(project_id, team_name)`` ловит точное совпадение. Процессная блокировка лишь снижает конкуренцию.
"""
import asyncio
import csv
import io
import json
import logging
import os
import shutil
import socket
import threading
import unicodedata
import zipfile
import zlib
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path, PurePosixPath
from typing import Any, NamedTuple, Optional
from urllib.parse import urlsplit

from fastapi import FastAPI, UploadFile
from sqlalchemy.exc import IntegrityError

from ..common import new_id, utc_now_iso
from ..db import session_scope
from ..repositories import submissions as repo
from ..state import get_storage
from . import projects
from .netguard import UnsafeUrl, validate_public_url
from .storage import ID_RE, PayloadTooLarge, SavedFile, Storage, UnsafePath, safe_join

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_BYTES = 2_000_000_000       # размер хранится в INTEGER (до 2^31 − 1)
MAX_FILES = 30
MAX_LINKS = 20
MAX_TEAM_NAME = 120
MAX_TABLE_BYTES = 5 * 1024 * 1024
MAX_TABLE_ROWS = 500
MAX_SCAN_ROWS = 10_000                  # сколько записей листа просматриваем, включая пустые
MAX_XLSX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_XLSX_ENTRIES = 1000
MAX_ZIP_ENTRIES = 65_534                # больше — ZIP64, такой архив не принимаем
READ_CHUNK = 1024 * 1024
URL_CHECK_WORKERS = 8

TABLE_COLUMNS = {"team_name": None, "repo_url": "repository", "host_url": "host",
                 "presentation_url": "presentation", "screencast_url": "screencast"}
REPOSITORY_HOSTS = ("github.com", "gitlab.com", "bitbucket.org")
EXTENSION_TYPES = {
    **dict.fromkeys((".pdf", ".pptx", ".ppt", ".key"), "presentation"),
    **dict.fromkeys((".mp4", ".mov", ".webm", ".mkv"), "screencast"),
    **dict.fromkeys((".zip", ".tar", ".gz", ".7z"), "archive"),
    **dict.fromkeys((".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".heic"), "image"),
}
SKIPPED_FOLDERS = ("__MACOSX",)

_LOCKS = tuple(threading.Lock() for _ in range(64))


class InvalidSubmission(ValueError):
    """Входные данные не годятся (422)."""


class DuplicateTeam(Exception):
    """Команда с таким названием в проекте уже есть (409)."""


class TableRow(NamedTuple):
    number: int
    fields: dict


class ArchiveSource(NamedTuple):
    saved: SavedFile
    folders: tuple


def _project_lock(project_id: str) -> threading.Lock:
    return _LOCKS[zlib.crc32(project_id.encode("utf-8")) % len(_LOCKS)]


# ─────────────────────────────── значения и проверки ───────────────────────────────

def clean_team_name(raw: Any) -> str:
    name = unicodedata.normalize("NFC", str(raw if raw is not None else "")).strip()
    if not name:
        raise InvalidSubmission("Название команды обязательно")
    if len(name) > MAX_TEAM_NAME:
        raise InvalidSubmission(f"Название команды длиннее {MAX_TEAM_NAME} символов")
    if any(unicodedata.category(ch) in ("Cc", "Cf") for ch in name):
        raise InvalidSubmission("В названии команды недопустимые символы")
    return name


def parse_links(raw: Optional[str]) -> list:
    if raw is None or not raw.strip():
        return []
    try:
        data = json.loads(raw)
    except ValueError:
        raise InvalidSubmission("links: ожидается JSON-массив ссылок") from None
    if not isinstance(data, list) or not all(isinstance(item, str) for item in data):
        raise InvalidSubmission("links: ожидается JSON-массив строк")
    cleaned = [item.strip() for item in data]
    if any(not item for item in cleaned):
        raise InvalidSubmission("links: пустая ссылка")
    cleaned = list(dict.fromkeys(cleaned))
    if len(cleaned) > MAX_LINKS:
        raise InvalidSubmission(f"Не больше {MAX_LINKS} ссылок")
    return cleaned


def file_material_type(filename: str) -> str:
    return EXTENSION_TYPES.get(PurePosixPath(filename).suffix.lower(), "other")


def link_material_type(url: str) -> str:
    parts = urlsplit(url)
    host = (parts.hostname or "").lower()
    if parts.path.lower().endswith(".git") or any(host == h or host.endswith("." + h) for h in REPOSITORY_HOSTS):
        return "repository"
    return "host"


def _resolver(app: FastAPI):
    return getattr(app.state, "url_resolver", None) or socket.getaddrinfo


def check_links(app: FastAPI, urls: list) -> list:
    """Нормализованные публичные адреса; первая небезопасная ссылка — UnsafeUrl."""
    resolver = _resolver(app)
    return [validate_public_url(url, resolver=resolver) for url in urls]


# ─────────────────────────────── файлы ───────────────────────────────

def _literal_dir(storage: Storage, *parts: str) -> Path:
    """Каталог внутри корня данных без симлинков на пути."""
    for part in parts:
        if not ID_RE.fullmatch(part) and part not in ("projects", "submissions", "raw"):
            raise UnsafePath("недопустимая часть пути")
    root = storage.root.resolve()
    resolved = safe_join(root, *parts)
    if resolved != root.joinpath(*parts):
        raise UnsafePath("симлинк в пути каталога")
    return resolved


def submission_dir(storage: Storage, project_id: str, submission_id: str) -> Path:
    return _literal_dir(storage, "projects", project_id, "submissions", submission_id)


def _relative(*parts: str) -> str:
    return "/".join(parts)


def _relocate(storage: Storage, project_id: str, submission_id: str, saved: SavedFile) -> str:
    """Переносит сохранённый файл в raw/<material_id>/ работы; возвращает относительный путь."""
    name = saved.path_rel.rsplit("/", 1)[-1]
    source_dir = _literal_dir(storage, "projects", project_id, "submissions", saved.material_id)
    target_dir = _literal_dir(storage, "projects", project_id, "submissions", submission_id, "raw", saved.material_id)
    target_dir.mkdir(parents=True, exist_ok=True)
    os.replace(source_dir / name, target_dir / name)
    try:
        source_dir.rmdir()
    except OSError:
        pass
    return _relative("projects", project_id, "submissions", submission_id, "raw", saved.material_id, name)


def _link_shared(storage: Storage, project_id: str, submission_id: str, shared: SavedFile, material_id: str) -> str:
    """Кладёт общий архив пачки в raw/ работы: жёсткая ссылка (без копии данных), иначе копия."""
    name = shared.path_rel.rsplit("/", 1)[-1]
    source = _literal_dir(storage, "projects", project_id, "submissions", shared.material_id) / name
    target_dir = _literal_dir(storage, "projects", project_id, "submissions", submission_id, "raw", material_id)
    target_dir.mkdir(parents=True, exist_ok=True)
    try:
        os.link(source, target_dir / name)
    except OSError:
        shutil.copyfile(source, target_dir / name)
    return _relative("projects", project_id, "submissions", submission_id, "raw", material_id, name)


def _discard(storage: Storage, project_id: str, submission_id: str, material_ids: list) -> None:
    """Best-effort: убирает всё, что успели записать для работы."""
    try:
        shutil.rmtree(submission_dir(storage, project_id, submission_id), ignore_errors=True)
        for material_id in material_ids:
            shutil.rmtree(_literal_dir(storage, "projects", project_id, "submissions", material_id), ignore_errors=True)
    except (UnsafePath, ValueError, OSError):
        logger.error("не удалось убрать файлы незавершённой работы %s", submission_id)


async def _chunks(upload: UploadFile):
    while True:
        chunk = await upload.read(READ_CHUNK)
        if not chunk:
            return
        yield chunk


async def _read_limited(upload: UploadFile, limit: int) -> bytes:
    if upload.size is not None and upload.size > limit:
        raise PayloadTooLarge(f"файл больше {limit} байт")
    parts, total = [], 0
    async for chunk in _chunks(upload):
        total += len(chunk)
        if total > limit:
            raise PayloadTooLarge(f"файл больше {limit} байт")
        parts.append(chunk)
    return b"".join(parts)


# ─────────────────────────────── запись в БД ───────────────────────────────

def summary_view(row: dict) -> dict:
    return {"id": row["id"], "team_name": row["team_name"], "intake_status": row["intake_status"],
            "completeness": row.get("completeness") or "unknown", "missing": list(row.get("missing") or []),
            "evaluation_status": "not_started", "flags": []}      # оценки и флаги подключаются при сборке


def material_view(row: dict) -> dict:
    ref = row.get("source_ref") or ""
    view = {"material_id": row["id"], "type": row["type"],
            "name": ref.rsplit("/", 1)[-1] if row["source_kind"] == "upload" else ref,
            "status": row["status"]}
    if row.get("reason") is not None:
        view["reason"] = row["reason"]
    if not view["name"]:
        del view["name"]
    return view


def _existing_keys(session, project_id: str) -> set:
    return {name.casefold() for name in repo.team_names(session, project_id)}


def ensure_project_and_name(app: FastAPI, project_id: str, team_name: Optional[str]) -> None:
    """LookupError — проекта нет; DuplicateTeam — быстрая проверка до записи файлов (окончательная — при вставке)."""
    with session_scope(app.state.session_factory) as session:
        projects.get_project(session, project_id)
        if team_name is not None and team_name.casefold() in _existing_keys(session, project_id):
            raise DuplicateTeam(team_name)


def _commit(app: FastAPI, project_id: str, submission_id: str, team_name: str, materials: list) -> dict:
    row = {"id": submission_id, "project_id": project_id, "team_name": team_name, "intake_status": "pending",
           "completeness": "unknown", "missing": None, "inventory_path": None, "version": 1,
           "created_at": utc_now_iso()}
    with _project_lock(project_id):         # внутри процесса; между процессами сериализует UPDATE проекта ниже
        try:
            with session_scope(app.state.session_factory) as session:
                # Первый же оператор транзакции — UPDATE строки проекта: он берёт блокировку записи (SQLite) или строки
                # (Postgres) до проверки названия, поэтому два создателя в разных процессах не проверяют его одновременно.
                projects.touch_project(session, project_id)
                project = projects.get_project(session, project_id)
                if team_name.casefold() in _existing_keys(session, project_id):
                    raise DuplicateTeam(team_name)
                repo.insert_submission(session, row)
                for material in materials:
                    repo.insert_material(session, {**material, "submission_id": submission_id})
                if project["status"] == "setup":
                    try:
                        projects.advance_status(session, project_id, "intake")
                    except ValueError:      # статус сменили параллельно — оставляем как есть
                        pass
        except IntegrityError:
            raise DuplicateTeam(team_name) from None
    return summary_view(row)


async def _create_one(app: FastAPI, project_id: str, team_name: str, *, uploads: list = (), links: list = (),
                      archive: Optional[SavedFile] = None) -> dict:
    """Сохраняет файлы и вставляет работу. ``links`` — пары (тип, нормализованный URL)."""
    storage = get_storage(app)
    submission_id = new_id()
    material_ids: list = []
    try:
        materials: list = []
        for upload in uploads:
            material_id = new_id()
            material_ids.append(material_id)
            saved = await storage.save_upload(project_id, "submissions", upload.filename or "file", _chunks(upload),
                                              max_bytes=MAX_FILE_BYTES, material_id=material_id)
            reference = await asyncio.to_thread(_relocate, storage, project_id, submission_id, saved)
            materials.append({"id": material_id, "type": file_material_type(saved.path_rel), "source_kind": "upload",
                              "source_ref": reference, "status": "available", "size_bytes": saved.size_bytes})
        if archive is not None:
            material_id = new_id()
            reference = await asyncio.to_thread(_link_shared, storage, project_id, submission_id, archive, material_id)
            materials.append({"id": material_id, "type": "archive", "source_kind": "upload", "source_ref": reference,
                              "status": "available", "size_bytes": archive.size_bytes})
        for material_type, url in links:
            materials.append({"id": new_id(), "type": material_type, "source_kind": "url", "source_ref": url,
                              "status": "available", "size_bytes": None})
        return await asyncio.to_thread(_commit, app, project_id, submission_id, team_name, materials)
    except BaseException:
        _discard(storage, project_id, submission_id, material_ids)
        raise


async def create_submission(app: FastAPI, project_id: str, team_name: str, uploads: list,
                            links_raw: Optional[str]) -> dict:
    """Одна работа. InvalidSubmission/UnsafeUrl — 422, PayloadTooLarge — 413, DuplicateTeam — 409, LookupError — 404."""
    name = clean_team_name(team_name)
    links = parse_links(links_raw)
    if len(uploads) > MAX_FILES:
        raise InvalidSubmission(f"Не больше {MAX_FILES} файлов")
    for upload in uploads:
        if upload.size is not None and upload.size > MAX_FILE_BYTES:
            raise PayloadTooLarge(f"файл больше {MAX_FILE_BYTES} байт")
    await asyncio.to_thread(ensure_project_and_name, app, project_id, name)
    normalized = await asyncio.to_thread(check_links, app, links)
    summary = await _create_one(app, project_id, name, uploads=uploads,
                                links=[(link_material_type(url), url) for url in normalized])
    return summary


# ─────────────────────────────── таблица пачки ───────────────────────────────

def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_csv(data: bytes) -> list:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise InvalidSubmission("Таблица должна быть в кодировке UTF-8") from None
    if "\x00" in text:
        raise InvalidSubmission("Таблица не похожа на CSV")
    first_line = text.split("\n", 1)[0]
    delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
    records = []
    try:
        for index, record in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)):
            if index >= MAX_SCAN_ROWS:
                raise PayloadTooLarge("в таблице слишком много строк")
            records.append([_normalize_cell(cell) for cell in record])
    except csv.Error:
        raise InvalidSubmission("Не удалось прочитать CSV") from None
    return records


def _read_xlsx(data: bytes) -> list:
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile:
        raise InvalidSubmission("Не удалось прочитать XLSX") from None
    if len(infos) > MAX_XLSX_ENTRIES or sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED:
        raise PayloadTooLarge("XLSX слишком велик после распаковки")
    from openpyxl import load_workbook
    records = []
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            for index, values in enumerate(workbook.worksheets[0].iter_rows(values_only=True)):
                if index >= MAX_SCAN_ROWS:
                    raise PayloadTooLarge("в таблице слишком много строк")
                records.append([_normalize_cell(cell) for cell in values])
        finally:
            workbook.close()
    except PayloadTooLarge:
        raise
    except Exception:       # openpyxl бросает разные типы на битых файлах; формулы не вычисляются, значения — текст
        raise InvalidSubmission("Не удалось прочитать XLSX") from None
    return records


def parse_table(data: bytes) -> list:
    """CSV (UTF-8, запятая или точка с запятой) или XLSX (первый лист) -> строки данных с номерами листа."""
    if not data.strip():
        raise InvalidSubmission("Таблица пуста")
    records = _read_xlsx(data) if data[:4] == b"PK\x03\x04" else _read_csv(data)
    if not records:
        raise InvalidSubmission("Таблица пуста")
    header = [cell.lower() for cell in records[0]]
    if "team_name" not in header:
        raise InvalidSubmission("В таблице нет колонки team_name")
    index = {name: header.index(name) for name in TABLE_COLUMNS if name in header}
    rows = [TableRow(number, {name: (record[i] if i < len(record) else "") for name, i in index.items()})
            for number, record in enumerate(records[1:], start=2)]
    while rows and not any(rows[-1].fields.values()):
        rows.pop()
    if len(rows) > MAX_TABLE_ROWS:
        raise PayloadTooLarge(f"в таблице больше {MAX_TABLE_ROWS} строк")
    return rows


# ─────────────────────────────── ZIP пачки ───────────────────────────────

def _zip_entry_count(path: Path) -> int:
    """Число записей из конца архива, не разбирая каталог целиком (защита от архива с миллионами записей)."""
    size = path.stat().st_size
    with open(path, "rb") as handle:
        handle.seek(max(0, size - 65_557))
        tail = handle.read()
    position = tail.rfind(b"PK\x05\x06")
    if position < 0 or len(tail) < position + 22:
        raise zipfile.BadZipFile("нет конца каталога")
    return int.from_bytes(tail[position + 10:position + 12], "little")


def _zip_folders(path: Path) -> tuple:
    """Папки верхнего уровня по каталогу архива, в порядке появления; содержимое не распаковывается."""
    if _zip_entry_count(path) > MAX_ZIP_ENTRIES:
        raise PayloadTooLarge("в архиве слишком много файлов")
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    folders: dict = {}
    for entry in names:
        parts = entry.replace("\\", "/").split("/")
        if len(parts) < 2 or not parts[0] or parts[0] in SKIPPED_FOLDERS or parts[0].startswith("."):
            continue
        folders.setdefault(parts[0], None)
    return tuple(folders)


async def _store_archive(app: FastAPI, project_id: str, upload: UploadFile) -> ArchiveSource:
    storage = get_storage(app)
    if upload.size is not None and upload.size > MAX_ARCHIVE_BYTES:
        raise PayloadTooLarge(f"файл больше {MAX_ARCHIVE_BYTES} байт")
    saved = await storage.save_upload(project_id, "submissions", upload.filename or "archive.zip", _chunks(upload),
                                      max_bytes=MAX_ARCHIVE_BYTES, material_id=new_id())
    try:
        path = _literal_dir(storage, "projects", project_id, "submissions", saved.material_id) / saved.path_rel.rsplit("/", 1)[-1]
        folders = await asyncio.to_thread(_zip_folders, path)
        if not folders:
            raise InvalidSubmission("В архиве нет папок команд")
        if len(folders) > MAX_TABLE_ROWS:
            raise PayloadTooLarge(f"в архиве больше {MAX_TABLE_ROWS} папок команд")
    except zipfile.BadZipFile:
        _drop_shared(storage, project_id, saved)
        raise InvalidSubmission("Архив повреждён или не является ZIP") from None
    except BaseException:
        _drop_shared(storage, project_id, saved)
        raise
    return ArchiveSource(saved, folders)


def _drop_shared(storage: Storage, project_id: str, saved: SavedFile) -> None:
    try:
        shutil.rmtree(_literal_dir(storage, "projects", project_id, "submissions", saved.material_id), ignore_errors=True)
    except (UnsafePath, ValueError, OSError):
        logger.error("не удалось убрать временный архив пачки")


# ─────────────────────────────── пачка ───────────────────────────────

class _Team(NamedTuple):
    label: str
    name: str
    links: tuple         # (тип, нормализованный URL)
    archive: bool


def _validate_urls(app: FastAPI, urls: list) -> dict:
    """URL -> нормализованный адрес или None (небезопасный/некорректный); проверки идут параллельно."""
    resolver = _resolver(app)

    def check(url: str) -> Optional[str]:
        try:
            return validate_public_url(url, resolver=resolver)
        except UnsafeUrl:
            return None

    unique = list(dict.fromkeys(urls))
    with ThreadPoolExecutor(max_workers=URL_CHECK_WORKERS) as pool:
        return dict(zip(unique, pool.map(check, unique)))


def _existing_keys_of(app: FastAPI, project_id: str) -> set:
    with session_scope(app.state.session_factory) as session:
        projects.get_project(session, project_id)
        return _existing_keys(session, project_id)


async def create_batch(app: FastAPI, project_id: str, table: Optional[UploadFile], archive: Optional[UploadFile]) -> dict:
    """Пачка работ. Плохие строки уходят в rejected; InvalidSubmission/PayloadTooLarge/LookupError — ошибка запроса."""
    if table is None and archive is None:
        raise InvalidSubmission("Нужна таблица (table) или архив (archive)")
    existing = await asyncio.to_thread(_existing_keys_of, app, project_id)      # LookupError — проекта нет
    rows: list = []
    if table is not None:
        rows = await asyncio.to_thread(parse_table, await _read_limited(table, MAX_TABLE_BYTES))

    storage = get_storage(app)
    source: Optional[ArchiveSource] = await _store_archive(app, project_id, archive) if archive is not None else None
    try:
        teams, rejected = await _plan_teams(app, rows, source, existing)
        items: list = []
        for team in teams:
            try:
                summary = await _create_one(app, project_id, team.name, links=team.links,
                                            archive=source.saved if (source and team.archive) else None)
            except DuplicateTeam:
                rejected.append((_order(team.label), team.label, "команда с таким названием уже существует"))
                continue
            except LookupError:
                raise
            except Exception as error:      # сбой одной команды не валит пачку; детали (пути) в лог не пишем
                logger.error("не удалось создать работу пачки (%s)", type(error).__name__)
                rejected.append((_order(team.label), team.label, "внутренняя ошибка при создании работы"))
                continue
            items.append(summary)
    finally:
        if source is not None:
            _drop_shared(storage, project_id, source.saved)
    rejected.sort(key=lambda entry: entry[0])
    return {"items": items, "rejected": [{"row": label, "reason": reason} for _, label, reason in rejected]}


def _order(label: str) -> tuple:
    return (0, int(label), "") if label.isdigit() else (1, 0, label)


async def _plan_teams(app: FastAPI, rows: list, source: Optional[ArchiveSource], existing: set) -> tuple:
    """Команды пачки (строки таблицы, затем папки архива без строки) и отклонённые строки.

    Отклонённая строка не занимает название: сначала проверяются ссылки, затем повторы внутри пачки."""
    rejected: list = []

    def reject(label: str, reason: str) -> None:
        rejected.append((_order(label), label, reason))

    candidates: list = []        # (строка, название, [(тип, URL, колонка)])
    for row in rows:
        label = str(row.number)
        if not any(row.fields.values()):
            reject(label, "пустая строка")
            continue
        try:
            name = clean_team_name(row.fields.get("team_name", ""))
        except InvalidSubmission as error:
            reject(label, str(error))
            continue
        candidates.append((label, name, [(TABLE_COLUMNS[column], row.fields[column], column)
                                         for column in TABLE_COLUMNS if TABLE_COLUMNS[column] and row.fields.get(column)]))

    checked = await asyncio.to_thread(_validate_urls, app, [url for _, _, links in candidates for _, url, _ in links])
    folder_keys: dict = {}       # нормализованное название -> исходное имя папки
    for folder in (source.folders if source else ()):
        try:
            folder_keys.setdefault(clean_team_name(folder).casefold(), folder)
        except InvalidSubmission as error:
            reject("".join(ch for ch in folder if ch.isprintable())[:MAX_TEAM_NAME] or "?", str(error))

    teams: list = []
    accepted: set = set()
    table_keys: set = set()
    for label, name, links in candidates:
        key = name.casefold()
        table_keys.add(key)
        bad = next((column for _, url, column in links if checked[url] is None), None)
        if bad is not None:
            reject(label, f"недопустимая ссылка в колонке {bad}")
        elif key in accepted:
            reject(label, "название команды повторяется в пачке")
        elif key in existing:
            reject(label, "команда с таким названием уже существует")
        else:
            accepted.add(key)
            teams.append(_Team(label, name, tuple((kind, checked[url]) for kind, url, _ in links), key in folder_keys))

    for key, folder in folder_keys.items():
        name = clean_team_name(folder)
        if key in accepted:
            continue                                  # папка привязана к принятой строке таблицы
        if key in table_keys:
            reject(name, "папка архива относится к отклонённой строке таблицы")
        elif key in existing:
            reject(name, "команда с таким названием уже существует")
        else:
            teams.append(_Team(name, name, (), True))
    if len(teams) > MAX_TABLE_ROWS:
        raise PayloadTooLarge(f"в пачке больше {MAX_TABLE_ROWS} команд")
    return teams, rejected


# ─────────────────────────────── чтение ───────────────────────────────

def list_submissions(session, project_id: str) -> list:
    """LookupError — проекта нет."""
    projects.get_project(session, project_id)
    return [summary_view(row) for row in repo.list_for_project(session, project_id)]


def get_submission(session, submission_id: str) -> dict:
    """LookupError — работы нет, она удалена или проект удалён."""
    row = repo.get(session, submission_id)
    if row is None:
        raise LookupError(f"работа {submission_id!r} не найдена")
    return {"id": row["id"], "team_name": row["team_name"],
            "materials": [material_view(m) for m in repo.materials_of(session, submission_id)],
            "total": None, "complete": row.get("completeness") == "complete",
            "criteria": [], "flags": [], "cost_usd": None, "latest_run_id": None}
